"""
Builds the same spreadsheet as create-sheet.gs, as an .xlsx file.

For when you would rather drag a file into Drive than run a script. Upload it,
then File → Save as Google Sheets. Everything survives the conversion except
protected ranges and cell notes, which only the Apps Script can set.

    python3 tools/sheet/build-xlsx.py

Writes "Sunday School Companion — Content.xlsx" beside itself.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "Sunday School Companion — Content.xlsx"

INK = "FF0E1731"
INK_TEXT = "FFF2F0EA"
SECTION = "FF16224A"
COMPUTED_BG = "FFF1F3F9"
COMPUTED_TEXT = "FF6B7280"
HELP_BG = "FFFBFAF7"

CLASSES = [
    ("6-7", "6–7 Years", "6–7 Years", 1, "Yes"),
    ("8-10", "8–10 Years", "8–10 Years", 2, "No"),
    ("11-13", "11–13 Years", "11–13 Years", 3, "No"),
    ("13-16", "13–16 Years", "13–16 Years", 4, "No"),
]

STATUSES = ["Draft", "Ready for Review", "Published"]
PRIORITIES = ["Core", "Supporting"]
GAME_IDEAS = [
    "Pick the right one",
    "Match things together",
    "Put things in order",
    "Explore and find",
]

# name, width (characters), wrap
CLASS_COLUMNS = [
    ("Chapter", 9, False),
    ("Title", 26, True),
    ("Bible Reference", 18, False),
    ("Curriculum", 62, True),
    ("Learning Objectives", 42, True),
    ("Memory Verse", 38, True),
    ("Memory Verse Reference", 20, False),
    ("Video", 26, False),
    ("Take Home", 38, True),
    ("Suggested Story Approach", 38, True),
    ("Suggested Game Approach", 30, True),
    ("Contributor", 18, False),
    ("Status", 18, False),
    ("Notes", 40, True),
]

LO_COLUMNS = [
    ("Class", 18, False),
    ("Chapter", 10, False),
    ("Objective ID", 14, False),
    ("Learning Objective", 62, True),
    ("Priority", 14, False),
    ("Notes", 40, True),
]

ROWS = 200


def header(ws, columns):
    for i, (name, width, wrap) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(color=INK_TEXT, bold=True, size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width

        if wrap:
            for r in range(2, ROWS + 2):
                ws.cell(row=r, column=i).alignment = Alignment(
                    vertical="top", wrap_text=True
                )
        else:
            for r in range(2, ROWS + 2):
                ws.cell(row=r, column=i).alignment = Alignment(vertical="top")

    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def dropdown(ws, column, values, strict):
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Pick one of the listed values."
    dv.showErrorMessage = strict
    ws.add_data_validation(dv)
    letter = get_column_letter(column)
    dv.add(f"{letter}2:{letter}{ROWS + 1}")


def computed(ws, column):
    for r in range(2, ROWS + 2):
        cell = ws.cell(row=r, column=column)
        cell.fill = PatternFill("solid", fgColor=COMPUTED_BG)
        cell.font = Font(color=COMPUTED_TEXT, italic=True, size=10)


def class_tab(wb, class_name):
    ws = wb.create_sheet(class_name)
    header(ws, CLASS_COLUMNS)
    ws.freeze_panes = "B2"

    dropdown(ws, 11, GAME_IDEAS, strict=False)
    dropdown(ws, 13, STATUSES, strict=True)

    # Objectives have one home. This column shows them, it does not store them.
    for r in range(2, ROWS + 2):
        ws.cell(row=r, column=5).value = (
            f'=IF($A{r}="","",IFERROR(TEXTJOIN(CHAR(10),TRUE,'
            f"FILTER('Learning Objectives'!$D$2:$D$500,"
            f"'Learning Objectives'!$A$2:$A$500=\"{class_name}\","
            f"'Learning Objectives'!$B$2:$B$500=$A{r})),\"\"))"
        )
    computed(ws, 5)

    for r in range(2, ROWS + 2):
        ws.cell(row=r, column=1).alignment = Alignment(
            horizontal="center", vertical="top"
        )
    return ws


def learning_objectives(wb):
    ws = wb.create_sheet("Learning Objectives")
    header(ws, LO_COLUMNS)

    dropdown(ws, 1, [c[2] for c in CLASSES], strict=True)
    dropdown(ws, 5, PRIORITIES, strict=True)

    for r in range(2, ROWS + 2):
        ws.cell(row=r, column=3).value = (
            f'=IF(OR($A{r}="",$B{r}=""),"",'
            f'TEXT($B{r},"00")&"."&COUNTIFS($A$2:$A{r},$A{r},$B$2:$B{r},$B{r}))'
        )
    computed(ws, 3)

    for r in range(2, ROWS + 2):
        for c in (2, 3, 5):
            ws.cell(row=r, column=c).alignment = Alignment(
                horizontal="center", vertical="top"
            )
    return ws


def config(wb):
    ws = wb.create_sheet("Config", 0)

    help_lines = [
        "Sunday School Companion — Content",
        "",
        "This sheet is where lessons begin. You write what a lesson is about; "
        "everything a child eventually sees is made from it later.",
        "",
        "To add a lesson:",
        "1.  Open your class tab.",
        "2.  Add a row. Fill in the chapter number, title and Bible reference.",
        "3.  Write the Curriculum — the lesson in your own words. This is the "
        "important one. Do not simplify it for children; that happens later.",
        "4.  Go to the Learning Objectives tab and add a row for each thing a "
        "child should come away with. One idea per row.",
        "5.  Add the memory verse, a video link and a take-home if the lesson "
        "has them. All optional.",
        '6.  Set Status to "Ready for Review". That is all you need to do.',
        "",
        "You never need to write story panels, dialogue, questions or answers. "
        "Those are made from what you write here.",
        "",
    ]

    for i, line in enumerate(help_lines, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1).font = Font(size=16, bold=True, color=INK)
    ws.cell(row=5, column=1).font = Font(bold=True)
    ws.cell(row=13, column=1).font = Font(italic=True, color=COMPUTED_TEXT)

    row = len(help_lines) + 1

    def table(row, title, headings, rows):
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, color=INK_TEXT)
        for c in range(1, len(headings) + 1):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=SECTION)
        row += 1
        for c, h in enumerate(headings, start=1):
            ws.cell(row=row, column=c, value=h).font = Font(
                bold=True, color=COMPUTED_TEXT
            )
        row += 1
        for r in rows:
            for c, v in enumerate(r, start=1):
                ws.cell(row=row, column=c, value=v).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
            row += 1
        return row + 1

    row = table(row, "Classes",
                ["Class ID", "Tab Name", "Display Name", "Order", "Live"],
                [list(c) for c in CLASSES])
    row = table(row, "Statuses", ["Status", "Means"], [
        ["Draft", "Being written. Not in the app."],
        ["Ready for Review", "Finished by the contributor. Not in the app yet."],
        ["Published", "Reviewed, made and approved. In the app. "
                      "Set by the production owner only."],
    ])
    row = table(row, "Objective Priority", ["Priority", "Means"], [
        ["Core", "Must be covered by a game."],
        ["Supporting", "Welcome, but optional."],
    ])
    table(row, "Game ideas a contributor can suggest",
          ["In the sheet", "What we build"], [
              ["Pick the right one", "Selection"],
              ["Match things together", "Pairing"],
              ["Put things in order", "Ordering"],
              ["Explore and find", "Discovery"],
              ["(blank)", "We decide"],
          ])

    for col, width in ((1, 46), (2, 26), (3, 26), (4, 10), (5, 52)):
        ws.column_dimensions[get_column_letter(col)].width = width
    for r in range(1, ws.max_row + 1):
        for c in range(1, 6):
            if ws.cell(row=r, column=c).fill.fgColor.rgb in (None, "00000000"):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=HELP_BG)
    ws.sheet_view.showGridLines = False
    return ws


def seed(sheets, lo):
    """The two chapters that exist in the repository. Curriculum is blank on
    both because neither ever had one — that is the gap this sheet closes."""
    ws = sheets["6–7 Years"]

    rows = [
        (1, "Baby Jesus at the Temple", "Luke 2:22–38", "",
         "My eyes have seen your salvation.", "Luke 2:30", "", "", "", "", "",
         "Draft",
         "Imported from content/baby-jesus-at-the-temple.story.json. No curriculum "
         "was ever written for it — the child-facing text was authored directly. "
         "CHECK: the repository says Luke 2:22–38, which includes Anna (vv. 36–38) "
         "and the chapter has two Anna panels; the brief said 22–33, which stops at "
         "Simeon. Art direction already agreed: cover is Mary holding Jesus with "
         "Simeon reaching out, warm temple light. \"Simeon held the baby\" is the "
         "heart of the chapter and gets the most space. It ends on gladness passed "
         "on, not on the ceremony finishing."),
        (2, "Stephen", "Acts 6–7", "",
         "Be kind to one another, forgiving one another.", "Ephesians 4:32", "",
         "", "", "", "", "Draft",
         "Imported from content/stephen.story.json. No curriculum was ever written "
         "for it. The memory verse translation is still PLACEHOLDER in the "
         "repository and currently raises a build warning — it needs a real "
         "translation. Art direction already agreed: we never show the stoning; "
         "the turning point is light from above with the crowd out of focus; "
         "aftermath, not event; the chapter must not end on grief."),
    ]

    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r[:4], start=1):
            ws.cell(row=i, column=c, value=v)
        for c, v in enumerate(r[4:], start=6):  # skip the formula in column E
            ws.cell(row=i, column=c, value=v)
        ws.row_dimensions[i].height = 130

    derived = ("Derived from a game the chapter already contains — not supplied "
               "by a contributor.")
    objectives = [
        ("6–7 Years", 1, "Recall who was waiting at the temple to see Jesus.",
         "Core", derived),
        ("6–7 Years", 2, "Recall how Stephen helped people.", "Core", derived),
        ("6–7 Years", 2, "Remember what Stephen prayed for the people who hurt him.",
         "Core", derived),
        ("6–7 Years", 2, "Understand the order the events happened in.",
         "Supporting", "Derived from the chapter’s existing sequence activity."),
    ]
    for i, o in enumerate(objectives, start=2):
        lo.cell(row=i, column=1, value=o[0])
        lo.cell(row=i, column=2, value=o[1])
        lo.cell(row=i, column=4, value=o[2])
        lo.cell(row=i, column=5, value=o[3])
        lo.cell(row=i, column=6, value=o[4])
        lo.row_dimensions[i].height = 46


def main():
    wb = Workbook()
    wb.remove(wb.active)

    config(wb)
    sheets = {c[1]: class_tab(wb, c[2]) for c in CLASSES}
    lo = learning_objectives(wb)
    seed(sheets, lo)

    wb.active = 1
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
